"""Data sources for labels: Excel, CSV/TSV, JSON or a pasted table, read into rows of text.

Every row becomes a label and every column a field. Cells come back as strings; date
columns are normalised to ISO dates (YYYY-MM-DD) so the label can format them as it
likes. Kept apart from the user interface, like the .qutie format.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import json
import re
from dataclasses import asdict, dataclass, field
from pathlib import Path

from .codes import valid_gs1

EXTENSIONS = (".xlsx", ".xlsm", ".csv", ".tsv", ".txt", ".json")
MAX_ROWS = 5000

YES = {"sì", "si", "yes", "y", "true", "vero", "x", "1", "ja", "oui", "sí", "sim", "да", "是", "はい"}
NO = {"no", "n", "false", "falso", "0", "nein", "non", "não", "нет", "否", "いいえ", ""}

DATE_PATTERNS = (
    (re.compile(r"^(\d{4})-(\d{1,2})-(\d{1,2})(?:[ T].*)?$"), "ymd"),
    (re.compile(r"^(\d{1,2})[/.\-](\d{1,2})[/.\-](\d{4})$"), "dmy"),
    (re.compile(r"^(\d{1,2})[/.\-](\d{1,2})[/.\-](\d{2})$"), "dmy2"),
)
NUMBER = re.compile(r"^[+-]?(\d{1,3}([.\s']\d{3})*|\d+)([.,]\d+)?$")


@dataclass
class Column:
    name: str
    type: str = "text"  # text | number | date | bool | ean13 | ean8


@dataclass
class Table:
    columns: list[Column] = field(default_factory=list)
    rows: list[list[str]] = field(default_factory=list)
    sheets: list[dict] = field(default_factory=list)  # [{"name", "rows"}] for workbooks
    sheet: str | None = None
    truncated: bool = False

    def as_dict(self) -> dict:
        return {"columns": [asdict(c) for c in self.columns], "rows": self.rows, "sheets": self.sheets,
                "sheet": self.sheet, "truncated": self.truncated}


# --------------------------------------------------------------------------- cells
def parse_date(text: str, day_first: bool = True) -> dt.date | None:
    text = text.strip()
    for pattern, kind in DATE_PATTERNS:
        m = pattern.match(text)
        if not m:
            continue
        a, b, c = (int(g) for g in m.groups()[:3])
        try:
            if kind == "ymd":
                return dt.date(a, b, c)
            year = c if kind == "dmy" else 2000 + c
            day, month = (a, b) if day_first or b > 12 else (b, a)
            return dt.date(year, month, day)
        except ValueError:
            return None
    return None


def _cell(value) -> str:
    """A workbook value as text: whole numbers without ".0", dates as ISO."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, dt.datetime):
        if (value.hour, value.minute, value.second) == (0, 0, 0):
            return value.date().isoformat()
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    return str(value).strip()


def detect_type(values: list[str]) -> str:
    """The type of a column, from its non-empty cells."""
    cells = [v.strip() for v in values if v and v.strip()]
    if not cells:
        return "text"
    if all(valid_gs1(c) and len(c) == 13 for c in cells):
        return "ean13"
    if all(valid_gs1(c) and len(c) == 8 for c in cells):
        return "ean8"
    if all(parse_date(c) for c in cells):
        return "date"
    if all(c.lower() in YES | NO for c in cells) and any(not c.isdigit() for c in cells):
        return "bool"
    if all(NUMBER.match(c) for c in cells):
        return "number"
    return "text"


def _finish(header: list[str] | None, body: list[list[str]], has_header: bool) -> Table:
    width = max([len(header or [])] + [len(r) for r in body]) if (header or body) else 0
    names = []
    for i in range(width):
        name = (header[i].strip() if header and i < len(header) and header[i] else "") if has_header else ""
        name = name or _column_letter(i)
        base, n = name, 2
        while name in names:  # field names must be unique
            name = f"{base} {n}"
            n += 1
        names.append(name)
    rows = [[(r[i] if i < len(r) else "") for i in range(width)] for r in body]
    rows = [r for r in rows if any(c.strip() for c in r)]
    truncated = len(rows) > MAX_ROWS
    rows = rows[:MAX_ROWS]
    columns = []
    for i, name in enumerate(names):
        kind = detect_type([r[i] for r in rows])
        if kind == "date":
            for r in rows:
                d = parse_date(r[i]) if r[i].strip() else None
                if d:
                    r[i] = d.isoformat()
        columns.append(Column(name, kind))
    return Table(columns, rows, truncated=truncated)


def _column_letter(i: int) -> str:
    s = ""
    i += 1
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


# --------------------------------------------------------------------------- readers
def parse_text(text: str, header: bool = True) -> Table:
    """CSV, TSV or cells copied from Excel / Google Sheets (tab separated)."""
    text = text.lstrip("﻿")
    lines = [line for line in text.splitlines() if line.strip()]
    if not lines:
        return Table()
    sample = "\n".join(lines[:20])
    counts = {sep: sample.count(sep) for sep in ("\t", ";", ",", "|")}
    sep = max(counts, key=counts.get) if max(counts.values()) else "\t"
    reader = csv.reader(io.StringIO("\n".join(lines)), delimiter=sep)
    data = [[c.strip() for c in row] for row in reader]
    return _finish(data[0] if header else None, data[1:] if header else data, header)


def _read_text_file(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", "replace")


def parse_json(text: str) -> Table:
    data = json.loads(text)
    if isinstance(data, dict):
        data = next((v for v in data.values() if isinstance(v, list)), [data])
    if not isinstance(data, list) or not data:
        return Table()
    if all(isinstance(r, dict) for r in data):
        names: list[str] = []
        for r in data:
            names.extend(k for k in r if k not in names)
        body = [[_cell(r.get(k)) if not isinstance(r.get(k), (dict, list)) else json.dumps(r.get(k), ensure_ascii=False)
                 for k in names] for r in data]
        return _finish(names, body, True)
    if all(isinstance(r, list) for r in data):
        rows = [[_cell(c) for c in r] for r in data]
        return _finish(rows[0], rows[1:], True)
    return _finish(["Value"], [[_cell(v)] for v in data], True)


def read_workbook(path: Path, sheet: str | None = None, header: bool = True) -> Table:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = []
        for ws in wb.worksheets:
            sheets.append({"name": ws.title, "rows": max(0, (ws.max_row or 1) - (1 if header else 0))})
        ws = wb[sheet] if sheet in wb.sheetnames else wb.worksheets[0]
        data = []
        for row in ws.iter_rows(values_only=True):
            data.append([_cell(v) for v in row])
            if len(data) > MAX_ROWS + 1:
                break
        while data and not any(data[-1]):
            data.pop()
        table = _finish(data[0] if header and data else None, data[1:] if header else data, header)
        # Sheets: the real count of rows with something in them, for the one we read.
        for s in sheets:
            if s["name"] == ws.title:
                s["rows"] = len(table.rows)
        table.sheets, table.sheet = sheets, ws.title
        return table
    finally:
        wb.close()


def read_file(path: str | Path, sheet: str | None = None, header: bool = True) -> Table:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        return read_workbook(path, sheet, header)
    if suffix == ".json":
        return parse_json(_read_text_file(path))
    if suffix == ".xls":
        raise ValueError("old .xls workbooks are not supported: save it as .xlsx")
    return parse_text(_read_text_file(path), header)


def describe(path: str | Path) -> dict:
    """Whether a linked file is still there, and when it changed."""
    p = Path(path)
    try:
        st = p.stat()
    except OSError:
        return {"exists": False, "path": str(p), "name": p.name}
    return {"exists": True, "path": str(p), "name": p.name, "size": st.st_size,
            "modified": dt.datetime.fromtimestamp(st.st_mtime).isoformat(timespec="seconds")}
